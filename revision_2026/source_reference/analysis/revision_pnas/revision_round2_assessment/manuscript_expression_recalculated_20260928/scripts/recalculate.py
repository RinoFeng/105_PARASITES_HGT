from pathlib import Path
from hashlib import sha256
import json, shutil
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT=Path(__file__).resolve().parents[1]
R=OUT.parent
SRC=R/'manuscript_expression_fonts_20260928'
BASE=R.parents[2]
LD=BASE/'analysis/Results_article/RNA-seq_expression/Leishmania_donovani'
for d in ['tables','data','checks','figures']:(OUT/d).mkdir(exist_ok=True)
inputs={}
def track(p):inputs[str(p)]=sha256(p.read_bytes()).hexdigest();return p
catalog=pd.read_excel(track(SRC/'tables/Table_S1_Catalogue.xlsx'),sheet_name='Catalogue_8989',dtype=str)
catalog=catalog[catalog.Specie.eq('Leishmania donovani')].copy()
mapping=pd.read_csv(track(LD/'genome_gff/gene_id2HGT_id.txt'),sep='\t',header=None,names=['reference_gene','reference_protein','reference_cds','Accession','identity_pct','alignment_length']).dropna(subset=['reference_gene','Accession'])
linked=mapping.merge(catalog[['hgt_id','Accession']],on='Accession',how='inner',validate='many_to_one')
linked['association_scope']='Existing reference-gene correspondence; not a new full-length sequence validation'
linked.to_csv(OUT/'data/Ld_catalogue_correspondence.tsv',sep='\t',index=False)
genes=sorted(linked.reference_gene.unique())
meta=pd.read_csv(track(OUT/'sources/PRJNA1220645_ENA_metadata.tsv'),sep='\t')
meta=meta[meta.run_accession.isin([f'SRR323122{i}' for i in range(77,83)])].copy()
meta['stage']=np.where(meta.experiment_title.str.contains('axenic amastigote'),'axenic_amastigote_like','promastigote')
meta['biological_replicate']=meta.experiment_title.str.extract(r'sample (\d+)').astype(int)
meta=meta.sort_values(['stage','biological_replicate'])
assert len(meta)==6 and meta.library_layout.eq('SINGLE').all()
assert meta.groupby('stage').size().to_dict()=={'axenic_amastigote_like':3,'promastigote':3}
order=meta.run_accession.to_list()
meta.to_csv(OUT/'data/Ld_sample_metadata.tsv',sep='\t',index=False)
matrices={};norm=[];sensitivity=[];gene_summary={}
for mode,filename,feature,strand in [('primary_reverse_gene','counts_s2.txt','gene','reverse'),('unstranded_gene','counts_s0.txt','gene','unstranded'),('legacy_unstranded_CDS','counts.txt','CDS','unstranded')]:
    f=track(LD/'counts'/filename);df=pd.read_csv(f,sep='\t',comment='#').set_index('Geneid')
    assert df.index.is_unique and df.Length.gt(0).all()
    samples=[c for c in df if c.endswith('.sorted.bam')]
    counts=df[samples].copy();counts.columns=counts.columns.str.replace('.sorted.bam','',regex=False)
    counts=counts[order];assert np.isfinite(counts).all().all() and counts.ge(0).all().all()
    rate=counts.div(df.Length,axis=0)
    denom=rate.sum();signal=rate.div(denom,axis=1)*1e6
    np.testing.assert_allclose(signal.sum(),1e6,rtol=1e-12)
    assert set(genes)<=set(signal.index)
    selected=signal.loc[genes]
    matrices[mode]=selected
    sig=selected.copy();sig.index.name='reference_gene';sig.to_csv(OUT/f'data/Ld_signal_{mode}.tsv',sep='\t',float_format='%.12g')
    for run in order:
        norm.append(dict(mode=mode,run_accession=run,feature=feature,strand=strand,denominator_gene_count=len(df),total_assigned_count=int(counts[run].sum()),sum_count_per_length=float(denom[run]),normalized_sum=float(signal[run].sum()),linked_genes=len(selected)))
    if mode=='primary_reverse_gene':
        source=counts.loc[genes].copy();source.insert(0,'feature_length_bp',df.loc[genes,'Length']);source.to_csv(OUT/'data/Ld_raw_counts_and_lengths.tsv',sep='\t')
        full=signal.copy();full.index.name='reference_gene';full.to_csv(OUT/'data/Ld_full_reference_signal.tsv.gz',sep='\t',compression='gzip',float_format='%.12g')
    for threshold in [1,5,10]:
        mask=selected.gt(threshold);stages={stage:meta.loc[meta.stage.eq(stage),'run_accession'].to_list() for stage in meta.stage.unique()}
        left=mask[stages['promastigote']];right=mask[stages['axenic_amastigote_like']]
        sensitivity.append(dict(mode=mode,threshold=threshold,linked_genes=len(selected),promastigote_all3=int(left.all(axis=1).sum()),amastigote_all3=int(right.all(axis=1).sum()),all6=int(mask.all(axis=1).sum()),both_stage_any=int((left.any(axis=1)&right.any(axis=1)).sum()),both_stage_2of3=int(((left.sum(axis=1)>=2)&(right.sum(axis=1)>=2)).sum()),both_stage_means=int(((selected[stages['promastigote']].mean(axis=1)>threshold)&(selected[stages['axenic_amastigote_like']].mean(axis=1)>threshold)).sum()),minimum_signal=float(selected.min().min())))
pd.DataFrame(norm).to_csv(OUT/'data/Ld_normalization_qc.tsv',sep='\t',index=False)
pd.DataFrame(sensitivity).to_csv(OUT/'data/Ld_feature_sensitivity.tsv',sep='\t',index=False)
primary=matrices['primary_reverse_gene']
long=primary.rename_axis('reference_gene').reset_index().melt(id_vars='reference_gene',var_name='run_accession',value_name='length_normalized_PROseq_signal').merge(meta[['run_accession','stage','biological_replicate']],on='run_accession',validate='many_to_one')
for threshold in [1,5,10]:long[f'signal_gt{threshold}']=long.length_normalized_PROseq_signal.gt(threshold)
long.to_csv(OUT/'data/Ld_signal_by_library.tsv',sep='\t',index=False)
summary=pd.DataFrame(sensitivity).query("mode == 'primary_reverse_gene'")
summary.to_csv(OUT/'data/Ld_threshold_summary.tsv',sep='\t',index=False)
fixed=['LDBPK_101360','LDBPK_150780','LDBPK_101370','LDBPK_171590']
primary.loc[fixed].to_csv(OUT/'data/FigureS8D_fixed_four_loci.tsv',sep='\t',float_format='%.12g')

# Independently rebuild the five-species counts from library-level values, retaining
# the existing correspondence eligibility rules and final-catalogue membership.
book=track(SRC/'tables/Table_S5_Expression.xlsx')
entry=pd.read_excel(book,sheet_name='Transcript_correspondence')
lib=pd.read_excel(book,sheet_name='TPM_by_library')
oldsummary=pd.read_excel(book,sheet_name='TPM_summary')
eligible=entry[entry.in_final_catalogue_same_species.eq(True)&entry.grade.ne('not_graded')].copy()
allfinal=entry[entry.in_final_catalogue_same_species.eq(True)].copy()
selected=lib.merge(eligible[['species','accession']],on=['species','accession'],how='inner',validate='many_to_one')
selected.TPM=pd.to_numeric(selected.TPM,errors='raise');assert selected.TPM.notna().all() and selected.TPM.ge(0).all()
assert not selected.duplicated(['species','accession','sample']).any()
maximum=selected.groupby(['species','accession']).TPM.max().rename('recomputed_max_TPM')
eligible=eligible.merge(maximum,on=['species','accession'],validate='one_to_one')
np.testing.assert_allclose(eligible.max_TPM,eligible.recomputed_max_TPM,rtol=1e-12)
recounts=[]
for species,e in eligible.groupby('species'):
    values=e.recomputed_max_TPM;sub=selected[selected.species.eq(species)]
    d={'species':species,'archived_entries':int(entry.species.eq(species).sum()),'final_catalogue_entries':int(allfinal.species.eq(species).sum()),'graded_entries':len(e),'not_graded':int(allfinal.species.eq(species).sum())-len(e),'libraries':sub['sample'].nunique()}
    for threshold in [1,5,10]:d[f'max_TPM_gt{threshold}']=int(values.gt(threshold).sum())
    d['at_least_two_libraries_TPM_gt10']=int(sub.assign(hit=sub.TPM.gt(10)).groupby('accession').hit.sum().ge(2).sum())
    for label,mask in [('0',values.eq(0)),('0<TPM<=1',values.gt(0)&values.le(1)),('1<TPM<=5',values.gt(1)&values.le(5)),('5<TPM<=10',values.gt(5)&values.le(10)),('TPM>10',values.gt(10))]:d[label]=int(mask.sum())
    recounts.append(d)
newsummary=pd.DataFrame(recounts)[oldsummary.columns]
newsummary.to_csv(OUT/'data/five_species_TPM_summary_recomputed.tsv',sep='\t',index=False)
eligible.to_csv(OUT/'data/five_species_maximum_TPM_recomputed.tsv',sep='\t',index=False)
diff=[]
for _,row in newsummary.iterrows():
    old=oldsummary[oldsummary.species.eq(row.species)].iloc[0]
    for key in newsummary.columns[1:]:
        if row[key]!=old[key]:diff.append({'species':row.species,'field':key,'previous':str(old[key]),'recomputed':str(row[key])})

for f in (SRC/'tables').glob('*.xlsx'):shutil.copy2(f,OUT/'tables'/f.name)
wb=openpyxl.load_workbook(book)
wb['Leishmania_stage_DE'].title='Leishmania_DE_legacy'
for row in wb['Contents'].iter_rows(min_row=2):
    if row[0].value=='Leishmania_stage_DE':
        row[0].value='Leishmania_DE_legacy';row[4].value='Archived four raw-count results. The dataset is PRO-seq; these estimates are not used as evidence of mature-RNA differential expression in the revised text.'
    if row[0].value=='TPM_summary':row[4].value='Recomputed from per-library TPM using existing catalogue and correspondence eligibility; strict TPM > thresholds.'
wb['README']['B3']='Five-species RNA-seq abundance and a separate six-library L. donovani PRO-seq signal analysis.'
wb['README']['B4']='TPM_summary and the Ld_* worksheets were recalculated. Original raw-count differential results are archived as Leishmania_DE_legacy, not interpreted as mature-RNA differential expression.'
wb['README'].append(['PRO-seq units','Length-normalized PRO-seq signal = 1e6*(gene count/feature length)/sum_all_reference_genes(count/length). TPM-scaled signal, not mature-transcript TPM.'])
wb['README'].append(['Sample identity','Public ENA metadata: SRR32312277–79 axenic amastigote-like; SRR32312280–82 promastigote. Original local README labels are reversed.'])
wb['README'].append(['Scope','27 reference genes are recoverable through the existing correspondence table and exact same-species catalogue accession matching. The former 35-gene list was not available; no claim that all 35 were reconstructed.'])
def sheet(name,df,description,replace=False):
    if name in wb:del wb[name]
    ws=wb.create_sheet(name);ws.append(list(df.columns))
    for row in df.itertuples(index=False,name=None):ws.append([None if pd.isna(v) else v for v in row])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='2F5496');cell.alignment=Alignment(wrap_text=True)
    for i,col in enumerate(df.columns,1):ws.column_dimensions[get_column_letter(i)].width=min(65,max(18,len(str(col))+2))
    if not replace:wb['Contents'].append([name,len(df),'Recalculated 2026-09-28; see sources and reproducible script','',description])
sheet('TPM_summary',newsummary,'Five-species RNA-seq counts recomputed from individual libraries.',True)
for name,df,desc in [('Ld_correspondence',linked,'Existing catalogue-to-reference gene links; source alignment fields retained.'),('Ld_signal_by_library',long,'Separate PRO-seq signal; not mature RNA TPM.'),('Ld_threshold_summary',summary,'Strict >1, >5, >10 signal; all-replicate and alternative stage summaries.'),('Ld_sample_metadata',meta,'Public run-to-experiment mapping and stages.'),('Ld_normalization_qc',pd.DataFrame(norm),'Normalization uses every reference gene in each count matrix, never only the HGT subset.'),('Ld_feature_sensitivity',pd.DataFrame(sensitivity),'Reverse-strand gene, unstranded gene and legacy CDS definitions.')]:sheet(name,df,desc)
wb.save(OUT/'tables/Table_S5_Expression.xlsx')
summary_out={'catalogue_Ldonovani_records':len(catalog),'existing_mapping_unique_genes':mapping.reference_gene.nunique(),'catalogue_linked_reference_genes':len(genes),'old_reported_35_members_reconstructed':False,'primary_all_six_gt10':int(summary.loc[summary.threshold.eq(10),'all6'].iloc[0]),'primary_minimum_signal':float(primary.min().min()),'ferrochelatase_signal':{k:float(v) for k,v in primary.loc['LDBPK_171590'].items()},'RNAseq_linked_entries':len(allfinal),'RNAseq_classifiable_entries':len(eligible),'RNAseq_gt1':int(newsummary.max_TPM_gt1.sum()),'RNAseq_gt5':int(newsummary.max_TPM_gt5.sum()),'RNAseq_gt10':int(newsummary.max_TPM_gt10.sum()),'prior_summary_differences':diff,'new_DE_test_performed':False,'input_hashes':inputs}
(OUT/'checks/recalculation_summary.json').write_text(json.dumps(summary_out,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in summary_out.items() if k!='input_hashes'},ensure_ascii=False,indent=2))
